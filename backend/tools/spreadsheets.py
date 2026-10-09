"""Create real downloadable Excel workbooks from structured sheets."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tools.artifacts import ARTIFACT_DIR, register_artifact, safe_filename


def create_xlsx(title: str, sheets: list[dict], filename: str | None = None) -> dict:
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("An Excel workbook needs at least one sheet.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for item in sheets[:20]:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "Sheet")[:31]
        sheet = workbook.create_sheet(name)
        rows = item.get("rows") or []
        if not isinstance(rows, list):
            continue
        for row in rows[:10000]:
            if isinstance(row, list):
                sheet.append([value if isinstance(value, (int, float, bool)) or value is None else str(value)[:32000] for value in row[:100]])
        if sheet.max_row:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2563EB")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 30
            sheet.sheet_view.showGridLines = False
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Aptos", size=11, color="17365D")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="EFF6FF")
                sheet.row_dimensions[row[0].row].height = 24
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(48, max(12, max((len(str(cell.value or "")) for cell in column[:100]), default=12) + 2))
                sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    if not workbook.sheetnames:
        raise ValueError("The workbook did not contain any valid sheets.")
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = safe_filename(filename or title, "xlsx")
    path = ARTIFACT_DIR / safe_name
    workbook.save(path)
    return register_artifact(path, safe_name, "xlsx")
