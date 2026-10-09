from __future__ import annotations
import io
import textwrap
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from tools.artifacts import ARTIFACT_DIR, register_artifact, safe_filename
from tools.visuals import select_visual


def create_pptx(title: str, slides: list[dict], filename: str | None = None, visual_assets: list[dict] | None = None) -> dict:
    if not isinstance(slides, list) or not slides:
        raise ValueError("A presentation needs at least one slide.")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    navy, blue, muted = RGBColor(20, 38, 65), RGBColor(36, 119, 220), RGBColor(90, 109, 131)
    assets = visual_assets or []
    used_images: set[str] = set()
    def text(slide, value, x, y, width, height, size=22, color=navy, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        for index, line in enumerate(str(value).split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = color
            paragraph.space_after = Pt(12)
        return box
    for index, data in enumerate(slides[:50]):
        if not isinstance(data, dict):
            continue
        content = data.get("content", data.get("bullets", []))
        if isinstance(content, str):
            content = content.splitlines()
        bullets = [chunk for item in content if isinstance(item, str) for chunk in textwrap.wrap(item, 260)] if isinstance(content, list) else []
        batches = [bullets[start:start + 4] for start in range(0, len(bullets), 4)] or [[]]
        for part, batch in enumerate(batches):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(247, 250, 254)
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.14), deck.slide_height)
            bar.fill.solid(); bar.fill.fore_color.rgb = blue; bar.line.fill.background()
            text(slide, str(title)[:100].upper(), .6, .28, 12, .35, 11, muted)
            heading = str(data.get("title") or f"Slide {index + 1}")[:150] + (" - continued" if part else "")
            text(slide, heading, .6, .82, 12, 1.05, 30, navy, True)
            asset = select_visual(assets, heading + " " + " ".join(batch) + (" " + title if index == 0 else ""), used_images) if not part else None
            width = 6.5 if asset else 11.9
            for row, bullet in enumerate(batch):
                text(slide, f"{row + 1:02}", .6, 2.1 + row * 1.1, .5, .4, 15, blue, True)
                text(slide, bullet, 1.25, 2.06 + row * 1.1, width - .65, 1.05, 17 if len(bullet) > 150 else 21)
            if asset:
                try:
                    picture = slide.shapes.add_picture(io.BytesIO(asset["data"]), Inches(8), Inches(2.1), width=Inches(4.65))
                    if picture.height > Inches(3.6):
                        ratio = Inches(3.6) / picture.height
                        picture.width = int(picture.width * ratio); picture.height = Inches(3.6)
                    caption = f"{asset.get('author', '')[:90]} / {asset.get('license', '')}\nImage source: Wikimedia Commons"
                    credit = text(slide, caption, 8, 5.9, 4.65, .7, 9, muted)
                    for paragraph in credit.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.hyperlink.address = asset.get("source", "")
                except (ValueError, OSError):
                    pass
            text(slide, "TEXENGINE  /  TEXDEV", .6, 7.05, 8, .25, 9, muted)
            text(slide, f"{len(deck.slides):02}", 12, 7.02, .7, .3, 11, muted)
        diagram = data.get("diagram", {})
        nodes = diagram.get("nodes", []) if isinstance(diagram, dict) else []
        if isinstance(nodes, list) and 2 <= len(nodes) <= 6 and all(isinstance(node, str) for node in nodes):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            text(slide, str(data.get("title", title))[:150], .6, .7, 12, 1.2, 30, navy, True)
            step = 12 / len(nodes)
            for node_index, node in enumerate(nodes):
                x = .6 + node_index * step
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.8), Inches(step - .35), Inches(1.8))
                shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(235, 244, 255)
                shape.line.color.rgb = blue
                text(slide, node[:60], x + .15, 3.12, step - .65, 1.2, 17, navy, True)
                if node_index < len(nodes) - 1:
                    text(slide, ">", x + step - .3, 3.35, .3, .4, 18, blue)
            text(slide, "TEXENGINE  /  TEXDEV", .6, 7.05, 8, .25, 9, muted)
    if not deck.slides:
        raise ValueError("No valid presentation slides were supplied.")
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = safe_filename(filename or title, "pptx")
    path = ARTIFACT_DIR / safe_name
    deck.save(path)
    if len(Presentation(path).slides) != len(deck.slides):
        raise ValueError("The generated PowerPoint could not be verified.")
    return {"success": True, "summary": {"title": title, "slides": len(deck.slides), "topics": [item.get("title", "") for item in slides if isinstance(item, dict)]}, **register_artifact(path, safe_name, "pptx")}
