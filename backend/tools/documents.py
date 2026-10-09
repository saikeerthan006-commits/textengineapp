from __future__ import annotations

import re
import io
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from tools.artifacts import ARTIFACT_DIR, register_artifact, safe_filename
from tools.visuals import select_visual


def _section_word_count(sections: list[dict]) -> int:
    text = []
    for section in sections:
        if not isinstance(section, dict):
            continue
        for value in (section.get("heading", ""), section.get("content", ""), *(section.get("paragraphs", []) if isinstance(section.get("paragraphs"), list) else []), *(section.get("bullet_points", section.get("bullets", [])) if isinstance(section.get("bullet_points", section.get("bullets", [])), list) else [])):
            if isinstance(value, str):
                text.append(value)
    return len(re.findall(r"\b[\w'-]+\b", " ".join(text)))



def _add_content(document, text):
    lines = text.splitlines()
    table_lines = [line.strip() for line in lines if line.strip().startswith("|")]
    if len(table_lines) >= 2 and re.search(r"\|\s*:?-{3,}", table_lines[1]):
        rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines if not re.fullmatch(r"[| :\-]+", line.strip())]
        if rows:
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for values in rows:
                cells = table.add_row().cells
                for index, value in enumerate(values[:len(cells)]):
                    cells[index].text = value.replace("**", "")
            return
    paragraph = document.add_paragraph()
    for index, part in enumerate(re.split(r"\*\*(.*?)\*\*", text, flags=re.S)):
        run = paragraph.add_run(part)
        run.bold = bool(index % 2)


def create_docx(title: str, sections: list[dict], filename: str | None = None, target_pages: int | None = None, visual_assets: list[dict] | None = None) -> dict:
    if not isinstance(sections, list) or not sections:
        raise ValueError("A Word document needs at least one section.")
    if target_pages is not None:
        target_pages = max(1, min(int(target_pages), 100))
        if len(sections) < target_pages:
            raise ValueError(f"The document needs {target_pages} content sections to honor the requested page count.")
        word_count = _section_word_count(sections)
        if word_count < target_pages * 250:
            raise ValueError(f"The document content is too short for {target_pages} pages ({word_count} words supplied).")
    document = Document()
    document.sections[0].top_margin = Inches(0.75)
    document.sections[0].bottom_margin = Inches(0.75)
    document.sections[0].left_margin = document.sections[0].right_margin = Inches(.85)
    normal = document.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        document.styles[style_name].font.color.rgb = RGBColor.from_string("17365D")
    document.styles["Title"].font.size = Pt(30)
    document.styles["Heading 1"].font.size = Pt(18)
    header = document.sections[0].header.paragraphs[0]
    header.text = "TEXENGINE  /  TEXDEV"
    header.style = document.styles["Caption"]
    document.sections[0].footer.paragraphs[0].text = str(title)[:100]
    document.add_heading(str(title or "Generated Document")[:200], 0)
    used_images: set[str] = set()
    for section_index, section in enumerate(sections[:100]):
        if not isinstance(section, dict):
            continue
        if target_pages and section_index:
            document.add_page_break()
        heading = str(section.get("heading", "Section"))[:200]
        document.add_heading(heading, level=1)
        paragraphs = section.get("paragraphs", [])
        content = section.get("content", "")
        if isinstance(content, str) and content.strip():
            paragraphs = [*paragraphs, *[part.strip() for part in re.split(r"\n\s*\n", content) if part.strip()]]
        for paragraph in paragraphs[:30] if isinstance(paragraphs, list) else []:
            if isinstance(paragraph, str) and paragraph.strip():
                _add_content(document, paragraph[:6000])
        asset = select_visual(visual_assets or [], heading + " " + str(content) + (" " + title if section_index == 0 else ""), used_images)
        if asset:
            try:
                picture = document.add_picture(io.BytesIO(asset["data"]), width=Inches(5.7))
                if picture.height > Inches(3.5):
                    picture.width = int(picture.width * Inches(3.5) / picture.height)
                    picture.height = Inches(3.5)
                document.add_paragraph(f"{asset.get('title', 'Illustration')} · {asset.get('author', '')} · {asset.get('license', '')}\n{asset.get('source', '')}", style="Caption")
            except (ValueError, OSError):
                pass
        bullets = section.get("bullet_points", section.get("bullets", []))
        for bullet in bullets[:30] if isinstance(bullets, list) else []:
            if isinstance(bullet, str) and bullet.strip():
                document.add_paragraph(bullet[:1000], style="List Bullet")
    for style_name in ("Normal", "Title", "Heading 1"):
        document.styles[style_name].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = safe_filename(filename or title, "docx")
    path = ARTIFACT_DIR / safe_name
    document.save(path)
    # Reopen the package so corrupt/empty output cannot be reported as success.
    verified = Document(path)
    if not verified.paragraphs:
        raise ValueError("The generated Word document could not be verified.")
    artifact = register_artifact(path, safe_name, "docx")
    return {"success": True, "summary": {"title": title, "topics": [item.get("heading", "") for item in sections if isinstance(item, dict)]}, **artifact}
