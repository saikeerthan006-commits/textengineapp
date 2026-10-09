"""Edit uploaded office files while retaining their native, editable format."""
import io
import uuid
from pathlib import Path
from docx import Document
from docx.shared import Pt as WordPt, RGBColor as WordColor
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from openpyxl import load_workbook
from tools.artifacts import ARTIFACT_DIR, register_artifact


def edit_office_file(filename: str, files: dict[str, bytes], replacements: list[dict] | None = None, theme: str = "preserve") -> dict:
    data = files.get(filename)
    if data is None:
        raise ValueError("Attach or select the office file to edit first.")
    replacements = replacements or []
    if not replacements and theme == "preserve":
        raise ValueError("Provide text replacements or choose the modern theme.")
    for item in replacements:
        if not isinstance(item, dict) or not isinstance(item.get("old"), str) or not item["old"] or not isinstance(item.get("new"), str):
            raise ValueError("Each replacement must contain nonempty old text and new text.")
    changed = 0
    def replace(text):
        nonlocal changed
        for item in replacements:
            changed += text.count(item["old"])
            text = text.replace(item["old"], item["new"])
        return text
    def paragraph_edit(paragraph):
        original = paragraph.text
        updated = replace(original)
        if original != updated and paragraph.runs:
            paragraph.runs[0].text = updated
            for run in paragraph.runs[1:]:
                run.text = ""
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        office = Document(io.BytesIO(data))
        def paragraphs(container):
            for paragraph in container.paragraphs:
                yield paragraph
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from paragraphs(cell)
        for paragraph in paragraphs(office):
            paragraph_edit(paragraph)
        if theme == "modern":
            office.styles["Normal"].font.name = "Aptos"
            office.styles["Normal"].font.size = WordPt(11)
            office.styles["Normal"].paragraph_format.line_spacing = 1.5
            for name in ("Title", "Heading 1", "Heading 2"):
                office.styles[name].font.color.rgb = WordColor.from_string("17365D")
    elif suffix == ".pptx":
        office = Presentation(io.BytesIO(data))
        for slide in office.slides:
            if theme == "modern":
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(247, 250, 254)
            for shape in slide.shapes:
                frames = [shape.text_frame] if shape.has_text_frame else []
                if shape.has_table:
                    frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
                for frame in frames:
                    for paragraph in frame.paragraphs:
                        paragraph_edit(paragraph)
                        if theme == "modern":
                            for run in paragraph.runs:
                                run.font.name = "Aptos"
                                run.font.color.rgb = RGBColor(20, 38, 65)
    elif suffix == ".xlsx":
        office = load_workbook(io.BytesIO(data))
        for sheet in office:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = replace(cell.value)
    else:
        raise ValueError("Office editing supports DOCX, PPTX, and XLSX.")
    if replacements and not changed:
        raise ValueError("None of the requested source text was found. Read the file and use its exact text.")
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    path = ARTIFACT_DIR / (uuid.uuid4().hex + suffix)
    try:
        office.save(path)
        result = register_artifact(path, Path(filename).stem + "_edited" + suffix, suffix[1:])
    finally:
        path.unlink(missing_ok=True)
    return {"success": True, "summary": {"replacements": changed, "theme": theme}, **result}
